"""客户名单导入（contracts/data-exchange.md §4；FR-029）。

- 支持 CSV（文本或文件）与 XLSX（openpyxl）；
- 匹配键：先 domain 后 email——命中已有公司则补联系方式，不新建；
- 导入前查退订黑名单：命中者联系方式标 invalid 并计入报告。
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from sqlalchemy.orm import Session, sessionmaker

from saler.core.compliance.suppression import is_suppressed
from saler.core.errors import ValidationFailed
from saler.core.leads import company_service, domain_service
from saler.core.leads.extractor import normalize_email
from saler.jobs.engine import JobContext
from saler.services.operation_log import log_op

HEADER_ALIASES: dict[str, set[str]] = {
    "email": {"email", "e-mail", "mail", "邮箱", "电子邮箱"},
    "phone": {"phone", "tel", "mobile", "whatsapp", "电话", "手机", "手机号"},
    "company_name": {"company_name", "company", "公司", "企业名称", "公司名称"},
    "domain": {"domain", "website", "web", "域名", "网站"},
    "country": {"country", "国家"},
    "language": {"language", "lang", "语言", "使用语言"},
    "tags": {"tags", "tag", "标签", "分类"},
    "note": {"note", "remark", "备注"},
}


def _map_headers(fieldnames: list[str]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for index, raw in enumerate(fieldnames):
        name = (raw or "").strip().lower()
        for canonical, aliases in HEADER_ALIASES.items():
            if name in aliases:
                mapping[canonical] = index
                break
    return mapping


def parse_csv_rows(text: str) -> list[dict]:
    reader = csv.reader(io.StringIO(text))
    try:
        header = next(reader)
    except StopIteration:
        return []
    mapping = _map_headers(list(header))
    if "email" not in mapping and "domain" not in mapping and "company_name" not in mapping:
        raise ValidationFailed("表头无法识别：至少需要 email/domain/company_name 之一（支持中英文表头）")
    rows: list[dict] = []
    for raw_row in reader:
        if not any(cell.strip() for cell in raw_row):
            continue
        rows.append(
            {
                canonical: (raw_row[idx].strip() if idx < len(raw_row) else "")
                for canonical, idx in mapping.items()
            }
        )
    return rows


def parse_xlsx_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = [str(cell) if cell is not None else "" for cell in next(iterator)]
    except StopIteration:
        return []
    mapping = _map_headers(header)
    if not mapping:
        raise ValidationFailed("表头无法识别（支持中英文表头）")
    rows: list[dict] = []
    for raw_row in iterator:
        cells = ["" if cell is None else str(cell).strip() for cell in raw_row]
        if not any(cells):
            continue
        rows.append(
            {
                canonical: (cells[idx] if idx < len(cells) else "")
                for canonical, idx in mapping.items()
            }
        )
    return rows


def load_rows(payload: dict) -> list[dict]:
    content = payload.get("content")
    path = payload.get("path")
    if content:
        return parse_csv_rows(content)
    if path:
        file_path = Path(path)
        if not file_path.exists():
            raise ValidationFailed(f"文件不存在：{path}")
        if file_path.suffix.lower() in (".xlsx", ".xlsm"):
            return parse_xlsx_rows(file_path)
        return parse_csv_rows(file_path.read_text(encoding="utf-8-sig"))
    raise ValidationFailed("需要提供 content（CSV 文本）或 path（文件路径）")


def _split_tags(raw: str) -> list[str]:
    return [tag.strip() for tag in raw.replace("；", ",").replace(";", ",").split(",") if tag.strip()]


def import_rows(db: Session, rows: list[dict]) -> dict:
    report = {
        "rows": len(rows),
        "new_companies": 0,
        "merged_companies": 0,
        "new_contacts": 0,
        "dup_contacts": 0,
        "suppressed": 0,
        "errors": [],
    }
    for index, row in enumerate(rows):
        email = normalize_email(row.get("email", "")) if row.get("email") else None
        phone = (row.get("phone") or "").strip()
        domain = (
            domain_service.normalize_domain(row["domain"]) if row.get("domain") else None
        )
        if not email and not phone and not domain and not row.get("company_name"):
            report["errors"].append({"index": index, "error": "整行为空或缺少可用字段"})
            continue
        try:
            company = None
            if domain:
                company = company_service.find_company_by_domain(db, domain)
            if company is None and email:
                company = company_service.find_company_by_email(db, email)
            if company is None:
                company = company_service.create_company_standalone(
                    db,
                    name=row.get("company_name") or domain or email,
                    country=(row.get("country") or "").lower()[:2] or None,
                    language=(row.get("language") or "").lower()[:10] or None,
                    source="import",
                )
                report["new_companies"] += 1
                if domain:
                    domain_row, _created = domain_service.get_or_create_domain(
                        db, domain=domain, source="import"
                    )
                    domain_row.company_id = company.id
                    db.commit()
            else:
                report["merged_companies"] += 1
                if row.get("company_name") and company.name in (None, domain, email):
                    company.name = row["company_name"]
                    db.commit()
            if row.get("tags"):
                merged = list(dict.fromkeys([*(company.tags or []), *_split_tags(row["tags"])]))
                company.tags = merged
                db.commit()

            if email:
                validity = "invalid" if is_suppressed(db, scope="email", value=email) else "valid"
                if validity == "invalid":
                    report["suppressed"] += 1
                _contact, created = company_service.upsert_contact(
                    db,
                    company_id=company.id,
                    type_="email",
                    value=row.get("email", "").strip(),
                    value_normalized=email,
                    category="general",
                    validity=validity,
                )
                report["new_contacts" if created else "dup_contacts"] += 1
            if phone:
                digits = "".join(ch for ch in phone if ch.isdigit())
                if 7 <= len(digits) <= 15:
                    _contact, created = company_service.upsert_contact(
                        db,
                        company_id=company.id,
                        type_="phone",
                        value=phone,
                        value_normalized=digits,
                        category="general",
                        validity="invalid"
                        if is_suppressed(db, scope="phone", value=digits)
                        else "valid",
                    )
                    report["new_contacts" if created else "dup_contacts"] += 1
        except Exception as exc:  # noqa: BLE001 - 导入报告式容错
            db.rollback()
            report["errors"].append({"index": index, "error": f"{type(exc).__name__}: {exc}"})
    return report


async def run_import(
    ctx: JobContext,
    session_factory: sessionmaker[Session],
    *,
    payload: dict,
) -> dict:
    ctx.progress(phase="parsing")
    rows = load_rows(payload)
    ctx.log(f"解析完成：{len(rows)} 行待导入")
    with session_factory() as db:
        report = import_rows(db, rows)
        log_op(
            db,
            "import",
            f"客户导入：新增公司 {report['new_companies']}，"
            f"新增联系方式 {report['new_contacts']}，黑名单命中 {report['suppressed']}",
            ref_type="job",
            ref_id=ctx.job_id,
            payload={k: v for k, v in report.items() if k != "errors"},
        )
    ctx.progress(phase="done", **{k: v for k, v in report.items() if k != "errors"})
    return report
