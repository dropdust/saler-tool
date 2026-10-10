"""线索导出（contracts/data-exchange.md §3；FR-028）。

- CSV：单表（买家）；XLSX：主表 + 被排除 + 统计 三个 Sheet（对齐既有工作习惯）；
- 被排除对象默认单独列出，不进买家主表。
"""

from __future__ import annotations

import csv
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy.orm import Session

from saler.core.errors import ValidationFailed
from saler.core.leads import company_service
from saler.db.models import Company, Contact

COLUMNS = [
    "company_name",
    "domain",
    "country",
    "language",
    "company_status",
    "tags",
    "contact_type",
    "contact_value",
    "contact_category",
    "contact_validity",
    "source_url",
    "excluded",
    "note",
]


def _rows_for(db: Session, companies: list[Company], *, excluded_flag: bool) -> list[dict]:
    rows: list[dict] = []
    for company in companies:
        domain_row = company_service.primary_domain(db, company.id)
        contacts = (
            db.query(Contact).filter_by(company_id=company.id).order_by(Contact.type, Contact.id).all()
        )
        base = {
            "company_name": company.name or "",
            "domain": domain_row.domain if domain_row else "",
            "country": company.country or "",
            "language": company.language or "",
            "company_status": company.status,
            "tags": ",".join(company.tags or []),
            "excluded": "是" if company.is_excluded else "",
            "note": company.note or "",
        }
        if not contacts:
            rows.append(
                {
                    **base,
                    "contact_type": "",
                    "contact_value": "",
                    "contact_category": "",
                    "contact_validity": "",
                    "source_url": "",
                }
            )
            continue
        for contact in contacts:
            rows.append(
                {
                    **base,
                    "contact_type": contact.type,
                    "contact_value": contact.value,
                    "contact_category": contact.category,
                    "contact_validity": contact.validity,
                    "source_url": contact.source_url or "",
                }
            )
    return rows


def _stats(rows: list[dict]) -> list[tuple[str, str, int]]:
    by_country: dict[str, int] = {}
    by_language: dict[str, int] = {}
    for row in rows:
        by_country[row["country"] or "(未标)" ] = by_country.get(row["country"] or "(未标)", 0) + 1
        by_language[row["language"] or "(未标)"] = by_language.get(row["language"] or "(未标)", 0) + 1
    stats = [("国家", code, count) for code, count in sorted(by_country.items())]
    stats += [("语言", code, count) for code, count in sorted(by_language.items())]
    return stats


def export_companies(
    db: Session,
    *,
    fmt: str,
    out_dir: Path,
    status: str | None = None,
    country: str | None = None,
) -> dict:
    if fmt not in ("csv", "xlsx"):
        raise ValidationFailed("format 必须为 csv 或 xlsx")
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")

    included = (
        company_service.list_companies(db, status=status, country=country, excluded=False).all()
    )
    excluded = company_service.list_companies(db, excluded=True).all()
    main_rows = _rows_for(db, included, excluded_flag=False)
    excluded_rows = _rows_for(db, excluded, excluded_flag=True)

    if fmt == "csv":
        path = out_dir / f"leads-{stamp}.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(main_rows)
        return {"path": str(path), "rows": len(main_rows), "excluded": len(excluded)}

    from openpyxl import Workbook

    path = out_dir / f"leads-{stamp}.xlsx"
    workbook = Workbook()
    buyers = workbook.active
    buyers.title = "买家"
    buyers.append(COLUMNS)
    for row in main_rows:
        buyers.append([row[column] for column in COLUMNS])
    excluded_sheet = workbook.create_sheet("被排除")
    excluded_sheet.append(COLUMNS)
    for row in excluded_rows:
        excluded_sheet.append([row[column] for column in COLUMNS])
    stats_sheet = workbook.create_sheet("统计")
    stats_sheet.append(["分组", "值", "行数"])
    for group, value, count in _stats(main_rows):
        stats_sheet.append([group, value, count])
    workbook.save(path)
    return {"path": str(path), "rows": len(main_rows), "excluded": len(excluded)}
